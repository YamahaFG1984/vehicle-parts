"""Read tables out of source files, keeping every cell's original value and exact location."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path


class ExtractionError(Exception):
    pass


@dataclass
class RawRow:
    locator: dict
    cells: list


@dataclass
class RawTable:
    name: str  # sheet title, or "p{page}t{n}" for PDF tables
    rows: list[RawRow] = field(default_factory=list)
    kind: str = "sheet"  # sheet / csv / pdf


@dataclass
class ExtractionResult:
    tables: list[RawTable]
    problems: list[dict] = field(default_factory=list)  # {"locator":…, "message":…}


def detect_file_type(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return "xlsx"
    if suffix in (".csv", ".tsv", ".txt"):
        return "csv"
    if suffix == ".pdf":
        return "pdf"
    hint = "；旧版 .xls 请先在 Excel 中另存为 .xlsx" if suffix == ".xls" else ""
    raise ExtractionError(f"不支持的文件类型：{path.suffix or '（无扩展名）'}（支持 .xlsx / .csv / .pdf{hint}）")


def extract(path: Path, file_type: str) -> ExtractionResult:
    return {"xlsx": extract_xlsx, "csv": extract_csv, "pdf": extract_pdf}[file_type](path)


def _is_blank(cells) -> bool:
    return all(c is None or (isinstance(c, str) and not c.strip()) for c in cells)


def extract_xlsx(path: Path) -> ExtractionResult:
    from openpyxl import load_workbook

    # Open by content (file handle), so a missing or odd extension does not matter.
    handle = path.open("rb")
    try:
        try:
            wb = load_workbook(handle, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for corrupt files
            raise ExtractionError(
                f"无法打开 Excel 文件（请确认是 .xlsx 格式；旧版 .xls 请先在 Excel 中另存为 .xlsx）：{exc}"
            ) from exc
        tables = []
        try:
            for ws in wb.worksheets:
                table = RawTable(name=ws.title)
                for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    cells = list(row)
                    if _is_blank(cells):
                        continue
                    table.rows.append(RawRow({"sheet": ws.title, "row": idx}, cells))
                tables.append(table)
        finally:
            wb.close()
    finally:
        handle.close()
    return ExtractionResult(tables)


def extract_csv(path: Path) -> ExtractionResult:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    table = RawTable(name=path.stem, kind="csv")
    for idx, row in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
        if _is_blank(row):
            continue
        table.rows.append(RawRow({"row": idx}, row))
    return ExtractionResult([table])


def extract_pdf(path: Path) -> ExtractionResult:
    import pdfplumber

    tables: list[RawTable] = []
    problems: list[dict] = []
    try:
        pdf = pdfplumber.open(path)
    except Exception as exc:
        raise ExtractionError(f"无法打开 PDF：{exc}") from exc
    with pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            found = page.extract_tables()
            if not found:
                found = page.extract_tables(
                    {"vertical_strategy": "text", "horizontal_strategy": "text"}
                )
            found = [t for t in found if t and len(t) >= 1 and max(len(r) for r in t) >= 3]
            if not found:
                text = (page.extract_text() or "").strip()
                problems.append({
                    "locator": {"page": page_no},
                    "message": "本页未找到表格" + ("（页面无文本层，可能是扫描件）" if not text else ""),
                })
                continue
            for t_no, rows in enumerate(found, start=1):
                table = RawTable(name=f"p{page_no}t{t_no}", kind="pdf")
                for r_no, row in enumerate(rows, start=1):
                    cells = [(c or "").replace("\n", " ").strip() for c in row]
                    if _is_blank(cells):
                        continue
                    table.rows.append(RawRow({"page": page_no, "table": t_no, "row": r_no}, cells))
                tables.append(table)
    return ExtractionResult(tables, problems)


def locator_label(file_name: str, locator: dict) -> str:
    if "sheet" in locator:
        return f"{file_name}!{locator['sheet']}!R{locator['row']}"
    if "page" in locator:
        tail = f"!t{locator['table']}!r{locator['row']}" if "table" in locator else ""
        return f"{file_name}!p{locator['page']}{tail}"
    return f"{file_name}!R{locator.get('row', '?')}"
