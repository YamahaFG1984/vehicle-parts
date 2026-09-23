"""Generate the incremental-import demo files in samples/demo/ (originals are only read)."""

from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from openpyxl import Workbook, load_workbook

SOURCE_A = "候选人材料_供应商A报价表.xlsx"
OUT_A = "供应商A报价表_v2.xlsx"
OUT_C = "供应商C目录.pdf"

# A v2 relative to the original: what changed and what the import should call it.
A_CHANGES = {
    "A-001": {"Unit Price": 44.10, "Quote Date": "2026-09-20"},  # UPDATED: new price
    "A-002": {"MOQ": 10},  # UPDATED: MOQ
    "A-003": {"Unit Price": 49.90, "Quote Date": "2026-09-20"},  # UPDATED: new price
    "A-006": {"Package Size": "151 x 91 x 20 cm"},  # UPDATED: new carton (soft evidence)
    "A-012": {"Truck Application": "Volvo VN 2006-2017"},  # CONFLICT: model years only overlap
}
A_REMOVED = {"A-027"}  # NOT_IN_LATEST
A_ADDED = [
    # same OE and attributes as A-005 → same-supplier suspect, never auto-merged
    ["A-028", "A-42-00", "Fan Shroud", "Volvo VNL 2018-2023", None, "OE-VNL-1005",
     "111 x 100 x 50 cm", 61.20, "USD", 5, "2026-09-20"],
    # no OE, same attributes as B-018 → cross-supplier suspect
    ["A-029", "A-43-R", "Side Grille, Right side", "Volvo VN 2004-2017", "Right", None,
     "65 x 55 x 36 cm", 139.00, "USD", 2, "2026-09-20"],
    # unknown category → needs data
    ["A-030", "A-44-00", "Hood Latch", "Volvo VNL 2018-2023", None, "OE-VNL-3001",
     "20 x 10 x 5 cm", 12.50, "USD", 20, "2026-09-20"],
]

C_HEADER = ["Part #", "Brand", "Description", "Application", "Side", "OEM #",
            "Carton Dimensions", "Price", "Currency", "MOQ", "Date"]
C_PAGE1 = [
    ["C-1001", "BrandX", "Front Grille", "Volvo VNL 2018-2023", "", "OE-VNL-1001",
     "1220 x 750 x 70 mm", "39.80", "USD", "3", "2026-09-15"],
    ["C-1002", "BrandX", "Bug Screen", "Volvo VNL 2018-2023", "", "OE-VNL-1002",
     "120 x 75 x 8 cm", "$44.00", "", "5", "2026-09-15"],
    ["C-1003-R", "BrandX", "Side Grille", "Volvo VNL 2018-2023", "Right", "OE-VNL-1003",
     "86 x 36 x 56 cm", "50.00", "USD", "2", "2026-09-15"],
    ["C-2003", "BrandY", "Air Cleaner Housing", "Freightliner Century", "", "OE-FRC-2003",
     "121 x 22 x 92 cm", "70.00", "CAD", "2", "2026-09-15"],
    ["C-2005", "BrandY", "Radiator Fan Shroud", "Freightliner Century", "", "OE-FRC-2005",
     "", "79.00", "USD", "2", "2026-09-15"],
    ["C-2006", "BrandY", "Front Grille", "Freightliner Century", "", "OE-FRC-2006",
     "124 x 77 x 10 cm", "92.50", "USD", "4", "09/16/2026"],
]
C_PAGE2 = [  # continued table, no header row on this page
    ["C-2004", "BrandY", "Air Filter Housing Cap", "Freightliner Century", "", "OE-FRC-2004",
     "32 x 32 x 8 cm", "75.00", "USD", "2", "2026-09-15"],
    ["C-9001", "BrandY", "Grille", "Kenworth T680 2014-2021", "", "",
     "130 x 80 x 12 cm", "88.00", "USD", "1", "2026-09-15"],
    ["C-1001", "BrandX", "Front Grille", "Volvo VNL 2018-2023", "", "OE-VNL-1001",
     "1220 x 750 x 70 mm", "38.90", "USD", "10", "2026-09-15"],
]


class Command(BaseCommand):
    help = "生成增量演示文件：samples/demo/供应商A报价表_v2.xlsx 与 供应商C目录.pdf"

    def add_arguments(self, parser):
        parser.add_argument("--out", default=str(settings.BASE_DIR / "samples" / "demo"))

    def handle(self, *args, **opts):
        out = Path(opts["out"])
        out.mkdir(parents=True, exist_ok=True)
        self.make_a_v2(settings.BASE_DIR / SOURCE_A, out / OUT_A)
        self.make_c_pdf(out / OUT_C)

    def make_a_v2(self, source: Path, target: Path):
        src = load_workbook(source, read_only=True)
        ws_src = src.worksheets[0]
        rows = [list(r) for r in ws_src.iter_rows(values_only=True)]
        title = ws_src.title
        src.close()
        header, body = rows[0], rows[1:]
        col = {h: i for i, h in enumerate(header)}
        wb = Workbook()
        ws = wb.active
        ws.title = title
        ws.append(header)
        for row in body:
            ref = row[col["Source Record ID"]]
            if ref in A_REMOVED:
                continue
            for field, value in A_CHANGES.get(ref, {}).items():
                row[col[field]] = value
            ws.append(row)
        for row in A_ADDED:
            ws.append(row)
        wb.save(target)
        self.stdout.write(self.style.SUCCESS(f"已生成 {target}"))

    def make_c_pdf(self, target: Path):
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import landscape, letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import (
                PageBreak,
                Paragraph,
                SimpleDocTemplate,
                Table,
                TableStyle,
            )
        except ImportError as exc:  # dev dependency only
            raise CommandError("生成 PDF 需要 reportlab（uv sync --group dev）") from exc
        style = TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
        ])
        styles = getSampleStyleSheet()
        doc = SimpleDocTemplate(str(target), pagesize=landscape(letter),
                                title="Supplier C Catalog (demo)")
        story = [
            Paragraph("Supplier C - Price Catalog 2026-09 (demo data)", styles["Title"]),
            Table([C_HEADER, *C_PAGE1], style=style, repeatRows=0),
            PageBreak(),
            Paragraph("continued", styles["Normal"]),
            Table(C_PAGE2, style=TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5)]),
                colWidths=None),
        ]
        doc.build(story)
        self.stdout.write(self.style.SUCCESS(f"已生成 {target}"))
