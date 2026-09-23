"""Write reviewer decisions from a filled-in review_list.xlsx back into the system (all or nothing)."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from apps.matching import services
from apps.matching.engine import run_matching
from apps.matching.models import Issue, MatchCandidate, ReviewDecision


class ReviewImportError(Exception):
    def __init__(self, errors: list[str]):
        super().__init__("\n".join(errors))
        self.errors = errors


@dataclass
class ReviewImportResult:
    candidates: int = 0
    issues: int = 0
    skipped: list[str] = field(default_factory=list)


def _rows(ws, id_col: str):
    headers = [str(c.value or "").strip() for c in ws[1]]
    for needed in (id_col, "decision", "note"):
        if needed not in headers:
            raise ReviewImportError([f"工作表 {ws.title} 缺少列 {needed}"])
    idx = {h: i for i, h in enumerate(headers)}
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        decision = str(row[idx["decision"]] or "").strip()
        if not decision:
            continue
        yield row_no, row[idx[id_col]], decision, str(row[idx["note"]] or "").strip()


def import_review(path: Path, reviewer: str, user=None) -> ReviewImportResult:
    wb = load_workbook(path, read_only=True, data_only=True)
    errors, plan, result = [], [], ReviewImportResult()
    open_status = {MatchCandidate.Status.PENDING, MatchCandidate.Status.NEEDS_INFO}
    if "疑似重复" in wb.sheetnames:
        for row_no, cid, decision, note in _rows(wb["疑似重复"], "candidate_id"):
            where = f"疑似重复!第{row_no}行"
            try:
                action = services.parse_action(decision)
                cand = MatchCandidate.objects.get(pk=int(cid))
            except (ValidationError, MatchCandidate.DoesNotExist, TypeError, ValueError) as exc:
                errors.append(f"{where}: {getattr(exc, 'message', exc)}")
                continue
            if action not in services.ACTION_TO_STATUS:
                errors.append(f"{where}: 候选只能填 merge / reject / needs_info")
                continue
            if cand.status not in open_status:
                result.skipped.append(f"{where}: 候选 #{cand.pk} 当前状态为 {cand.get_status_display()}，已跳过")
                continue
            hard = [c["field"] for c in cand.conflicts if c.get("hard")]
            if action == ReviewDecision.Action.MERGE and hard and not note:
                errors.append(f"{where}: 候选 #{cand.pk} 有硬冲突（{', '.join(hard)}），merge 必须填写 note")
                continue
            plan.append(("candidate", cand, action, note))
    if "异常清单" in wb.sheetnames:
        for row_no, iid, decision, note in _rows(wb["异常清单"], "issue_id"):
            where = f"异常清单!第{row_no}行"
            try:
                action = services.parse_action(decision)
                issue = Issue.objects.get(pk=int(iid))
            except (ValidationError, Issue.DoesNotExist, TypeError, ValueError) as exc:
                errors.append(f"{where}: {getattr(exc, 'message', exc)}")
                continue
            if action not in (ReviewDecision.Action.RESOLVE, ReviewDecision.Action.IGNORE):
                errors.append(f"{where}: 异常只能填 resolve / ignore")
                continue
            plan.append(("issue", issue, action, note))
    wb.close()
    if errors:
        raise ReviewImportError(errors)
    with transaction.atomic():
        for kind, obj, action, note in plan:
            if kind == "candidate":
                services.decide_candidate(obj, action, user=user, actor_name=reviewer, note=note,
                                          via=ReviewDecision.Via.XLSX, recluster=False)
                result.candidates += 1
            else:
                services.decide_issue(obj, action, user=user, actor_name=reviewer, note=note,
                                      via=ReviewDecision.Via.XLSX)
                result.issues += 1
        if result.candidates:
            run_matching()
    return result
