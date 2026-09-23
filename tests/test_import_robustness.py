"""Unseen formats: failed header detection, retry after fixing the mapping, reprocessing,
and part-number spelling changes."""

import pytest
from openpyxl import Workbook, load_workbook

from apps.catalog.models import SupplierItem
from apps.ingestion.mapping import MappingOverride
from apps.ingestion.models import ImportBatch, SourceFile
from apps.ingestion.services import ImportFailed, import_source
from apps.matching.models import Issue

from .conftest import FILE_A

pytestmark = pytest.mark.django_db

ODD_HEADERS = ["Artikel", "Bezeichnung", "Fahrzeug", "Karton", "Preis", "Waehrung"]
ODD_ROWS = [
    ["X-100", "Front Grille", "Volvo VNL 2018-2023", "122 x 75 x 7 cm", 40.5, "USD"],
    ["X-101", "Bug Screen", "Volvo VNL 2018-2023", "120 x 75 x 8 cm", 30.0, "USD"],
]
ODD_MAPPING = {"Artikel": "supplier_part_no", "Bezeichnung": "name", "Fahrzeug": "fitment",
               "Karton": "package_dims", "Preis": "price", "Waehrung": "currency"}


@pytest.fixture
def odd_file(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Lieferant X Preisliste"])  # title row above the header
    ws.append(ODD_HEADERS)
    for row in ODD_ROWS:
        ws.append(row)
    path = tmp_path / "lieferant_x.xlsx"
    wb.save(path)
    return path


def test_unrecognized_headers_fail_without_archiving(odd_file):
    with pytest.raises(ImportFailed) as exc:
        import_source(odd_file, "X")
    assert "column_aliases" in str(exc.value) and "--mapping" in str(exc.value)
    assert not SourceFile.objects.exists()
    assert not ImportBatch.objects.exists()


def test_retry_with_mapping_after_failure(odd_file):
    with pytest.raises(ImportFailed):
        import_source(odd_file, "X")
    result = import_source(odd_file, "X", override=MappingOverride(columns=ODD_MAPPING))
    assert result.batch.stats["new"] == 2
    assert SupplierItem.objects.get(supplier_part_no="X-100").dims_cm == [122.0, 75.0, 7.0]


def test_forced_header_row_exposes_columns_for_mapping(odd_file):
    preview = import_source(odd_file, "X", dry_run=True,
                            override=MappingOverride(header_row=2))
    assert preview.parse.mappings[0]["headers"] == ODD_HEADERS
    assert preview.parse.mappings[0]["columns"] == {}  # now the user can map them
    assert not preview.parse.unmatched


def test_dry_run_reports_unmatched_table_samples(odd_file):
    preview = import_source(odd_file, "X", dry_run=True)
    assert preview.parse.unmatched[0]["rows"][1]["cells"] == ODD_HEADERS


def test_archived_file_without_rows_can_be_retried(odd_file):
    # Files archived by older versions with a zero-row batch must not be stuck as "duplicate".
    import_source(odd_file, "X", override=MappingOverride(columns=ODD_MAPPING))
    ImportBatch.objects.update(stats={"rows": 0})
    result = import_source(odd_file, "X", override=MappingOverride(columns=ODD_MAPPING))
    assert not result.duplicate
    assert SourceFile.objects.count() == 1
    assert result.batch.mapping["reprocessed"] is True


def test_reprocess_successful_file(odd_file):
    import_source(odd_file, "X", override=MappingOverride(columns=ODD_MAPPING))
    assert import_source(odd_file, "X").duplicate
    result = import_source(odd_file, "X", override=MappingOverride(columns=ODD_MAPPING),
                           reprocess=True)
    assert result.batch.stats["unchanged"] == 2
    assert SourceFile.objects.count() == 1


def test_part_number_spelling_change_is_same_item(tmp_path):
    import_source(FILE_A, "A")
    wb = load_workbook(FILE_A)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[1].value == "A-03-L":
            row[1].value = "A03L"
    changed = tmp_path / "a_respelled.xlsx"
    wb.save(changed)
    result = import_source(changed, "A")
    assert result.batch.stats["not_in_latest"] == 0
    assert result.batch.stats.get("new", 0) == 0
    item = SupplierItem.objects.get(supplier__code="A", source_ref="A-003")
    assert item.supplier_part_no == "A03L"
    assert Issue.objects.filter(item=item, code="PART_NO_FORMAT_CHANGED").exists()
    assert item.field_values.filter(field="supplier_part_no", raw_value="A-03-L",
                                    is_current=False).exists()


def test_reprocess_applies_corrected_mapping(odd_file):
    # First import misses the price column: no offer, MISSING_PRICE open.
    partial = {k: v for k, v in ODD_MAPPING.items() if k not in ("Preis", "Waehrung")}
    import_source(odd_file, "X", override=MappingOverride(columns=partial))
    item = SupplierItem.objects.get(supplier_part_no="X-100")
    assert not item.offers.exists()
    assert Issue.objects.filter(item=item, code="MISSING_PRICE", status="open").exists()
    # Same file, same raw rows, corrected mapping: must be re-applied, not skipped as unchanged.
    result = import_source(odd_file, "X", override=MappingOverride(columns=ODD_MAPPING),
                           reprocess=True)
    assert result.batch.stats.get("unchanged", 0) == 0
    assert result.batch.stats["updated"] == 2
    offer = item.offers.get(is_current=True)
    assert (str(offer.price), offer.currency) == ("40.5000", "USD")
    assert not Issue.objects.filter(item=item, code="MISSING_PRICE", status="open").exists()


def test_reprocess_with_same_mapping_is_unchanged(odd_file):
    import_source(odd_file, "X", override=MappingOverride(columns=ODD_MAPPING))
    result = import_source(odd_file, "X", override=MappingOverride(columns=ODD_MAPPING),
                           reprocess=True)
    assert result.batch.stats["unchanged"] == 2


def test_new_synonym_reaches_imported_data(tmp_path, settings):
    """Maintenance path from the docs: add a phrase to synonyms.yaml, then run_matching."""
    import shutil

    import yaml

    from apps.core import rules
    from apps.matching.engine import run_matching

    wb = Workbook()
    ws = wb.active
    ws.append(["Part No", "Description", "Fitment", "Package", "Price", "Currency"])
    ws.append(["E-1", "Hood Emblem", "Volvo VNL 2018-2023", "20 x 10 x 3 cm", 9.5, "USD"])
    path = tmp_path / "emblems.xlsx"
    wb.save(path)
    import_source(path, "E")
    run_matching()
    item = SupplierItem.objects.get(supplier_part_no="E-1")
    assert item.category == ""
    assert Issue.objects.filter(item=item, code="UNKNOWN_CATEGORY", status="open").exists()

    rules_dir = tmp_path / "rules"
    shutil.copytree(settings.RULES_DIR, rules_dir)
    syn = yaml.safe_load((rules_dir / "synonyms.yaml").read_text(encoding="utf-8"))
    syn["categories"]["hood_emblem"] = {"label": "Hood Emblem", "sided": False,
                                        "phrases": ["hood emblem"]}
    (rules_dir / "synonyms.yaml").write_text(yaml.safe_dump(syn, allow_unicode=True),
                                             encoding="utf-8")
    settings.RULES_DIR = rules_dir
    try:
        summary = run_matching()
        item.refresh_from_db()
        assert item.category == "hood_emblem"
        assert summary.renormalized == 1
        assert not Issue.objects.filter(item=item, code="UNKNOWN_CATEGORY", status="open").exists()
        history = item.field_values.filter(field="name").order_by("created")
        assert [fv.normalized_value for fv in history] == [None, "hood_emblem"]
        assert [fv.is_current for fv in history] == [False, True]
        assert run_matching().renormalized == 0  # idempotent
    finally:
        rules.reload()
