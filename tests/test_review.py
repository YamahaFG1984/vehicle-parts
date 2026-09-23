"""Human decisions: priority over rules, reopening on new evidence, audit, xlsx write-back."""

import pytest
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from apps.exports.exporters import export_review
from apps.exports.review_import import ReviewImportError, import_review
from apps.ingestion.services import import_source
from apps.matching import services
from apps.matching.engine import run_matching
from apps.matching.models import MatchCandidate, ReviewDecision

from .conftest import DEMO_DIR

pytestmark = pytest.mark.django_db
MC = MatchCandidate
A = ReviewDecision.Action


def test_merge_suspect_joins_products(imported, pair, item, reviewer):
    cand = pair("A-001", "A-013") or pair("B-001", "A-013")
    services.decide_candidate(cand, A.MERGE, user=reviewer, note="同一产品，重复录入")
    assert item("A-013").product_id == item("A-001").product_id == item("B-001").product_id
    assert cand.decisions.filter(action=A.MERGE, actor=reviewer).exists()


def test_reject_survives_rerun(imported, pair, reviewer):
    cand = pair("A-021", "B-019")
    services.decide_candidate(cand, A.REJECT, user=reviewer, note="OE 录错")
    run_matching()
    cand.refresh_from_db()
    assert (cand.status, cand.decided_by) == (MC.Status.REJECTED, MC.DecidedBy.HUMAN)


def test_merge_over_hard_conflict_requires_note(imported, pair, reviewer):
    cand = pair("A-023", "B-021")
    with pytest.raises(ValidationError):
        services.decide_candidate(cand, A.MERGE, user=reviewer, note="")


def test_split_an_auto_merge(imported, pair, item, reviewer):
    cand = pair("A-005", "B-005")
    services.decide_candidate(cand, A.REJECT, user=reviewer, note="图片显示不同")
    assert item("A-005").product_id != item("B-005").product_id
    run_matching()
    assert item("A-005").product_id != item("B-005").product_id


def test_changed_evidence_reopens_human_decision(imported, pair, reviewer):
    cand = pair("A-006", "B-006")
    services.decide_candidate(cand, A.REJECT, user=reviewer, note="暂不合并")
    import_source(DEMO_DIR / "供应商A报价表_v2.xlsx", "A")  # A-006 dims change
    run_matching()
    cand.refresh_from_db()
    assert cand.decided_by == MC.DecidedBy.RULE
    assert cand.status == MC.Status.PENDING
    assert cand.decisions.filter(action=A.REOPEN).exists()


def test_xlsx_round_trip(imported, pair, item, tmp_path):
    path = export_review(tmp_path / "review.xlsx")
    wb = load_workbook(path)
    ws = wb["疑似重复"]
    headers = [c.value for c in ws[1]]
    target = pair("A-021", "B-019").pk
    for row in ws.iter_rows(min_row=2):
        if row[headers.index("candidate_id")].value == target:
            row[headers.index("decision")].value = "reject"
            row[headers.index("note")].value = "OE 共用，类型不同"
    wb.save(path)
    result = import_review(path, reviewer="张复核")
    assert result.candidates == 1
    cand = pair("A-021", "B-019")
    assert cand.status == MC.Status.REJECTED
    assert cand.decisions.get(action=A.REJECT).via == ReviewDecision.Via.XLSX


def test_xlsx_errors_write_nothing(imported, pair, tmp_path):
    path = export_review(tmp_path / "review.xlsx")
    wb = load_workbook(path)
    ws = wb["疑似重复"]
    headers = [c.value for c in ws[1]]
    ws.cell(row=2, column=headers.index("decision") + 1, value="reject")
    ws.cell(row=3, column=headers.index("decision") + 1, value="maybe")
    wb.save(path)
    with pytest.raises(ReviewImportError) as exc:
        import_review(path, reviewer="x")
    assert "第3行" in exc.value.errors[0]
    assert not ReviewDecision.objects.filter(via=ReviewDecision.Via.XLSX).exists()
