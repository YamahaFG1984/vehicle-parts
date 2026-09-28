"""SKU images and stock snapshots."""

import io
from datetime import date

import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from PIL import Image

from apps.catalog import images, stock
from apps.catalog.models import ItemImage, StockLevel

pytestmark = pytest.mark.django_db


def png(name="photo.png", size=(1200, 900), color=(242, 183, 5)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return SimpleUploadedFile(name, buf.getvalue(), content_type="image/png")


@pytest.fixture
def staff(client, django_user_model):
    user = django_user_model.objects.create_user("staff", password="pw-12345-x", is_staff=True)
    client.force_login(user)
    return user


# ---------------------------------------------------------------------------- images


def test_upload_several_images_first_is_primary(client, staff, imported, item):
    a001 = item("A-001")
    resp = client.post(f"/items/{a001.pk}/images/", {
        "images": [png("front.png"), png("side.png")], "source": "own_photo", "caption": "实拍"})
    assert resp.status_code == 302
    imgs = list(a001.images.order_by("created"))
    assert [i.is_primary for i in imgs] == [True, False]
    assert imgs[0].width == 1200 and imgs[0].source == "own_photo" and imgs[0].uploaded_by == staff
    thumb = Image.open(imgs[0].thumbnail.open("rb"))
    assert max(thumb.size) <= 480
    served = client.get(imgs[0].thumb_url)
    assert served.status_code == 200 and served["Content-Type"] == "image/jpeg"


@pytest.mark.parametrize(("upload", "message"), [
    (SimpleUploadedFile("fake.png", b"not really an image"), "不是有效的图片"),
    (SimpleUploadedFile("anim.gif", b"GIF89a"), "只接受"),
])
def test_rejects_non_images(imported, item, upload, message):
    with pytest.raises(ValidationError, match=message):
        images.add_images(item("A-001"), [upload])
    assert not ItemImage.objects.exists()


def test_rejects_oversized(imported, item, monkeypatch):
    monkeypatch.setattr(images, "MAX_BYTES", 100)
    with pytest.raises(ValidationError, match="超过"):
        images.add_images(item("A-001"), [png()])


def test_one_bad_file_stores_nothing(imported, item):
    with pytest.raises(ValidationError):
        images.add_images(item("A-001"), [png(), SimpleUploadedFile("x.png", b"junk")])
    assert not ItemImage.objects.exists()


def test_primary_switch_and_delete(imported, item):
    first, second = images.add_images(item("A-001"), [png("1.png"), png("2.png")])
    images.set_primary(second)
    first.refresh_from_db()
    assert not first.is_primary and ItemImage.objects.get(pk=second.pk).is_primary
    path = second.image.path
    images.delete_image(ItemImage.objects.get(pk=second.pk))
    first.refresh_from_db()
    assert first.is_primary  # the remaining image takes over
    import os
    assert not os.path.exists(path)


def test_images_follow_their_sku_across_merge_and_split(client, staff, imported, item, pair):
    from apps.matching import services

    images.add_images(item("A-001"), [png("a.png")])
    images.add_images(item("B-001"), [png("b.png")])
    product = item("A-001").product
    html = client.get(f"/products/{product.code}/").content.decode()
    assert html.count('class="gallery-item') == 2  # both SKUs' images on the product page
    services.decide_candidate(pair("A-001", "B-001"), "reject", user=staff, note="拆分")
    a_page = client.get(f"/products/{item('A-001').product.code}/").content.decode()
    assert a_page.count('class="gallery-item') == 1


def test_image_requires_login(client, imported, item):
    img = images.add_images(item("A-001"), [png()])[0]
    assert client.get(img.get_absolute_url()).status_code == 302


# ---------------------------------------------------------------------------- stock


def test_manual_stock_and_history(client, staff, imported, item):
    a001 = item("A-001")
    for qty, day in [(10, "2026-09-01"), (7, "2026-09-20")]:
        resp = client.post(f"/items/{a001.pk}/stock/", {
            "warehouse": "主仓", "quantity": qty, "as_of": day, "note": ""})
        assert resp.status_code == 302
    current = a001.stock_levels.get(is_current=True)
    assert (current.quantity, current.as_of, current.recorded_by) == (7, date(2026, 9, 20), staff)
    assert a001.stock_levels.filter(is_current=False, quantity=10).exists()
    assert stock.set_stock(a001, 7, as_of=date(2026, 9, 20)) is None  # unchanged


def test_negative_stock_rejected(client, staff, imported, item):
    resp = client.post(f"/items/{item('A-001').pk}/stock/", {
        "warehouse": "主仓", "quantity": -1, "as_of": "2026-09-20"})
    assert resp.status_code == 200 and not StockLevel.objects.exists()


def test_product_page_and_search_show_stock(client, staff, imported, item):
    stock.set_stock(item("A-001"), 5, warehouse="主仓")
    stock.set_stock(item("A-001"), 2, warehouse="芝加哥仓")
    stock.set_stock(item("B-001"), 4)
    product = item("A-001").product
    resp = client.get(f"/products/{product.code}/")
    assert resp.context["stock_total"] == 11
    result = client.get("/search/", {"q": "OE-VNL-1001"}).context["results"][0]
    assert result["stock_total"] == 11


def _stock_file(tmp_path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    path = tmp_path / "erp_stock.xlsx"
    wb.save(path)
    return path


def test_import_stock_file(imported, item, tmp_path):
    path = _stock_file(tmp_path, {
        # Two sheets with different column orders: each row must use its own sheet's columns.
        "主仓": [["库存导出"], ["供应商", "SKU", "仓库", "库存数量", "截至日期"],
               ["A", "a0100", "主仓", 12, "2026-09-25"],
               ["B", "B01X", "主仓", "3", "2026-09-25"],
               ["A", "ZZZ-1", "主仓", 1, "2026-09-25"],
               ["A", "A-02-00", "主仓", "abc", "2026-09-25"]],
        "外仓": [["qty", "part no", "warehouse"],
               [8, "A-03-L", "Dallas"]],
    })
    result = stock.import_stock(path, as_of=date(2026, 9, 26))
    assert dict(result.stats) == {"updated": 3, "unmatched": 1, "invalid": 1}
    a001 = item("A-001").stock_levels.get(is_current=True)
    assert (a001.quantity, a001.as_of, a001.source) == (12, date(2026, 9, 25), "import")
    assert a001.locator == "erp_stock.xlsx!主仓!R3"
    dallas = item("A-003").stock_levels.get(is_current=True)
    assert (dallas.warehouse, dallas.quantity, dallas.as_of) == ("Dallas", 8, date(2026, 9, 26))
    assert result.unmatched[0]["part_no"] == "ZZZ-1"
    assert "数量无法识别" in result.problems[0]["message"]


def test_import_stock_needs_columns(imported, tmp_path):
    path = _stock_file(tmp_path, {"S": [["foo", "bar"], [1, 2]]})
    with pytest.raises(stock.StockImportError, match="编号"):
        stock.import_stock(path)


def test_web_stock_import(client, staff, imported, item, tmp_path):
    path = _stock_file(tmp_path, {"S": [["SKU", "Qty"], ["A-01-00", 9]]})
    resp = client.post("/stock/import/", {
        "file": SimpleUploadedFile("erp.xlsx", path.read_bytes()), "supplier": "A",
        "warehouse": "", "as_of": ""}, follow=True)
    assert resp.status_code == 200
    assert item("A-001").stock_levels.get(is_current=True).quantity == 9
    assert "库存导入 #" in resp.content.decode()


def test_exports_include_stock_and_images(imported, item):
    from apps.exports.exporters import export_master, export_search, scope_from_search

    stock.set_stock(item("A-001"), 6)
    images.add_images(item("A-001"), [png()])
    for fn, args, members_sheet in [(export_master, (), "成员条目"),
                                    (export_search, (scope_from_search("A-01-00"),), "品牌号")]:
        buf = io.BytesIO()
        fn(buf, *args)
        buf.seek(0)
        wb = load_workbook(buf, read_only=True)
        stock_rows = list(wb["库存"].iter_rows(min_row=2, values_only=True))
        assert ("A", "A-01-00", "主仓", 6) in [r[1:5] for r in stock_rows]
        header = next(wb[members_sheet].iter_rows(max_row=1, values_only=True))
        row = next(r for r in wb[members_sheet].iter_rows(min_row=2, values_only=True)
                   if "A-01-00" in r)
        assert row[header.index("库存")] == 6 and row[header.index("图片数")] == 1


@pytest.mark.django_db(transaction=True)  # TRUNCATE cannot run inside the per-test transaction
def test_reset_removes_images_and_stock(imported, item):
    import os

    from apps.core.reset import reset_business_data

    img = images.add_images(item("A-001"), [png()])[0]
    stock.set_stock(item("A-001"), 1)
    path = img.image.path
    reset_business_data()
    assert not ItemImage.objects.exists() and not StockLevel.objects.exists()
    assert not os.path.exists(path)


# ---------------------------------------------------------------------------- FilePond uploads


@pytest.mark.parametrize("url_name", ["item", "imports"])
def test_upload_pages_use_filepond(client, staff, imported, item, url_name):
    url = f"/items/{item('A-001').pk}/" if url_name == "item" else "/imports/"
    html = client.get(url).content.decode()
    assert "vendor/filepond/filepond.min.js" in html and "js/uploads.js" in html
    assert "https://" not in html  # vendored, no CDN
    if url_name == "item":
        assert 'data-filepond' in html and 'data-upload-image-max="2000"' in html
        assert 'data-upload-accept=".jpg,.jpeg,.png,.webp"' in html and "multiple" in html
    else:
        assert html.count("data-filepond") == 2  # document import + stock import
        assert 'data-upload-accept=".xlsx,.xlsm,.csv,.tsv,.txt,.pdf"' in html
        assert 'data-upload-accept=".xlsx,.xlsm,.csv"' in html


def test_other_pages_do_not_load_filepond(client, staff, imported):
    assert "filepond" not in client.get("/").content.decode()


def test_vendored_filepond_files_present():
    from pathlib import Path

    root = Path(__file__).resolve().parent.parent / "static" / "vendor" / "filepond"
    versions = (root / "VERSIONS.txt").read_text()
    for name in ["filepond.min.js", "filepond.min.css", "filepond-plugin-image-resize.min.js",
                 "filepond-plugin-image-transform.min.js", "LICENSE.txt"]:
        assert (root / name).stat().st_size > 0, name
    assert "4.32.12" in versions
