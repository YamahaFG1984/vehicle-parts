"""Search → export: the export contains exactly what the search shows, plus everything
needed to check it (all brand numbers, per-supplier offers, provenance, related review items)."""

import io

import pytest
from django.core.management import call_command
from openpyxl import load_workbook

from apps.catalog.selectors import search
from apps.exports.exporters import ExportScope, export_master, export_search, scope_from_search
from apps.exports.review_import import import_review
from apps.matching.models import MatchCandidate

pytestmark = pytest.mark.django_db


def _book(scope):
    buf = io.BytesIO()
    export_search(buf, scope)
    buf.seek(0)
    return load_workbook(buf, read_only=True)


def _rows(wb, sheet):
    rows = list(wb[sheet].iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, r, strict=False)) for r in rows[1:] if any(v is not None for v in r)]


def test_export_matches_search_results(imported):
    for criteria in [("OE-VNL-1001", "", ""), ("grille", "", ""), ("", "B", ""),
                     ("Volvo", "", "已确认归一")]:
        scope = scope_from_search(*criteria)
        shown = {r["product"].pk for r in search(*criteria[:1], supplier=criteria[1],
                                                 state=criteria[2], limit=None)}
        assert scope.product_ids == shown
        wb = _book(scope)
        codes = {r["产品编号"] for r in _rows(wb, "归一产品")}
        assert len(codes) == len(shown)


def test_number_search_exports_whole_normalized_product(imported, item):
    wb = _book(scope_from_search("B01X"))
    assert [r["产品编号"] for r in _rows(wb, "归一产品")] == [item("A-001").product.code]
    members = _rows(wb, "品牌号")
    # The hit is B01X, but the product's other brand number (A-01-00) is exported too.
    assert {(m["供应商编号"], m["命中查询"]) for m in members} == {("B01X", "命中"), ("A-01-00", None)}
    offers = _rows(wb, "供应商报价")
    assert {(o["供应商"], o["价格"], o["币种"]) for o in offers} == {("A", 42.67, "USD"),
                                                                ("B", 46.58, "USD")}
    fields = {(f["供应商编号"], f["字段"]) for f in _rows(wb, "字段来源")}
    assert ("A-01-00", "package_dims") in fields and ("B01X", "price") in fields
    criteria = {r["项目"]: r["值"] for r in _rows(wb, "查询条件")}
    assert criteria["关键词 / 编号"] == "B01X" and criteria["命中归一产品数"] == 1


def test_review_items_are_limited_to_the_results(imported, pair):
    wb = _book(scope_from_search("OE-SHARED-01"))
    cand_ids = {r["candidate_id"] for r in _rows(wb, "疑似重复")}
    assert pair("A-021", "B-019").pk in cand_ids
    assert pair("A-023", "B-021").pk not in cand_ids  # unrelated products stay out
    issue_items = {r["供应商编号"] for r in _rows(wb, "异常清单")}
    assert issue_items <= {"A-27-00", "B28X", "A-01-00", "B01X", "A-14-00", "A-02-00", "B02X",
                           "A-16-00", "B15X"}


def test_search_export_can_be_written_back(imported, pair, tmp_path):
    path = tmp_path / "result.xlsx"
    export_search(path, scope_from_search("OE-SHARED-01"))
    wb = load_workbook(path)
    ws = wb["疑似重复"]
    headers = [c.value for c in ws[1]]
    target = pair("A-021", "B-019").pk
    for row in ws.iter_rows(min_row=2):
        if row[headers.index("candidate_id")].value == target:
            row[headers.index("decision")].value = "reject"
    wb.save(path)
    assert import_review(path, reviewer="复核人").candidates == 1
    assert pair("A-021", "B-019").status == MatchCandidate.Status.REJECTED


def test_empty_search_exports_criteria_only(imported):
    wb = _book(scope_from_search("no-such-part-xyz"))
    assert _rows(wb, "归一产品") == []
    assert {r["项目"]: r["值"] for r in _rows(wb, "查询条件")}["命中归一产品数"] == 0


def test_full_export_is_unscoped(imported):
    buf = io.BytesIO()
    export_master(buf, ExportScope())
    buf.seek(0)
    wb = load_workbook(buf, read_only=True)
    assert len(_rows(wb, "成员条目")) == 53
    assert "命中查询" not in [c.value for c in next(wb["成员条目"].iter_rows(max_row=1))]


def test_web_search_export(client, django_user_model, imported):
    client.force_login(django_user_model.objects.create_user("u", password="pw-12345-x"))
    page = client.get("/search/", {"q": "OE-VNL-1001"})
    assert "/exports/search/?q=OE-VNL-1001" in page.content.decode()
    resp = client.get("/exports/search/", {"q": "OE-VNL-1001"})
    assert resp.status_code == 200
    assert "attachment" in resp["Content-Disposition"]
    wb = load_workbook(io.BytesIO(b"".join(resp.streaming_content)), read_only=True)
    assert wb.sheetnames[:5] == ["查询条件", "说明", "归一产品", "品牌号", "供应商报价"]


def test_cli_search_export(imported, tmp_path):
    call_command("export_data", "--query", "OE-VNL-1001", "--out", str(tmp_path))
    files = list(tmp_path.glob("查询结果_OE-VNL-1001_*.xlsx"))
    assert len(files) == 1
    wb = load_workbook(files[0], read_only=True)
    assert len(_rows(wb, "归一产品")) == 1
